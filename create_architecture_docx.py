import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_color):
    """Sets cell background shading color in hex format (e.g. '1F4E78')."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="D3D3D3"):
    """Sets subtle light gray borders for table cells."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_architecture_doc():
    doc = docx.Document()

    # Page Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Global Style
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x26, 0x26, 0x26)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("AI-Swing-Trade Platform")
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)  # Deep Navy Blue

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(24)
    run_sub = p_sub.add_run("System Architecture Specification & Technical Blueprint")
    run_sub.font.size = Pt(14)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(8)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
        return h

    def add_bullet(bold_prefix="", text=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        if bold_prefix:
            r1 = p.add_run(bold_prefix)
            r1.font.bold = True
        if text:
            p.add_run(text)
        return p

    # Section 1: Executive Summary & System Scope
    add_h1("1. Executive System Overview")
    doc.add_paragraph(
        "The AI-Swing-Trade platform is an autonomous, cloud-native quantitative trading ecosystem built specifically for "
        "the Indian Equity market (NSE). The system integrates market screening, state-graph workflow orchestration, "
        "institutional risk budgeting, regulatory tax sentinel checks, Google Sheets persistence, Telegram bot supervision, "
        "and live broker API execution via Dhan HQ."
    )

    # Technology Stack Table
    add_h2("Technology Stack & Subsystems")
    t_tech = doc.add_table(rows=7, cols=3)
    t_tech.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_tech)

    tech_headers = ["Layer / Subsystem", "Technology / Framework", "Core Responsibility"]
    for i, title in enumerate(tech_headers):
        cell = t_tech.rows[0].cells[i]
        cell.text = title
        set_cell_background(cell, "1F4E78")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.size = Pt(9.5)

    tech_data = [
        ["User Interface & Master Dashboard", "Streamlit 1.30+ & Plotly", "Interactive strategy control, charts, telemetry, and manual trade triggers."],
        ["Workflow Orchestration", "LangGraph & Python 3.10+", "State-graph execution pipeline (Sync -> Screen -> Risk -> Execute -> Format)."],
        ["Market Data & Technical Analysis", "yFinance & pandas / numpy", "Real-time & historical OHLCV data, EMA, SMA, RSI, ATR calculations."],
        ["Broker Integration & Gateway", "Dhan HQ API v2 & Requests", "OAuth token validation, live holdings, funds limits, order routing."],
        ["Persistence Layer", "Google Sheets API (gspread)", "Cloud storage for Account metrics, OpenPositions, TradeHistory, and PaperTrades."],
        ["Automated Bot Supervisor", "python-telegram-bot & AsyncIO", "Background daemons for schedule alerts, scans, and Telegram bot commands."]
    ]

    for row_idx, row_data in enumerate(tech_data, start=1):
        row_cells = t_tech.rows[row_idx].cells
        bg_color = "F2F4F7" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = cell_value
            set_cell_background(row_cells[col_idx], bg_color)
            p = row_cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx < 2 else WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.size = Pt(9)
                if col_idx == 0:
                    run.font.bold = True

    # Section 2: Architectural Layers & Microservice Topology
    add_h1("2. System Microservice Architecture")
    doc.add_paragraph(
        "The system employs a Consolidated 3-Tier Microservices Model designed to operate within cloud resource limits "
        "while providing 0ms cold-start trade execution and zero single-point-of-failure isolation."
    )

    add_h2("Consolidated Microservice Tier Breakdown")
    add_bullet("1. Service Alpha (swing-trading-master): ", "Houses Streamlit UI, 4 Quantitative Strategy Engines, Backtesting Sandbox, and Master Telegram Supervisor Daemon.")
    add_bullet("2. Service Beta (nexus-dhan-gateway): ", "Dedicated Dhan HQ Broker REST Proxy (Port 8085). Caches holdings, cash margin limits, and order placement with a 30s RAM TTL to eliminate Dhan 429 rate limit issues.")
    add_bullet("3. Service Gamma (nexus-intelligence-api): ", "Market Intelligence REST Proxy (Port 8086). Serves News RSS Sentiment scores, FII/DII net flows, and stock fundamental ratios with a 5-minute RAM TTL.")

    # Microservice Endpoints Table
    add_h2("Microservices API Endpoint Registry")
    t_api = doc.add_table(rows=8, cols=4)
    t_api.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_api)

    api_headers = ["Microservice", "Endpoint", "Method", "Primary Responsibility"]
    for i, title in enumerate(api_headers):
        cell = t_api.rows[0].cells[i]
        cell.text = title
        set_cell_background(cell, "2E75B6")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.size = Pt(9.5)

    api_data = [
        ["Service Beta (Broker)", "/api/v1/health", "GET", "Health check, Dhan connection state, and Google Sheets status."],
        ["Service Beta (Broker)", "/api/v1/funds", "GET", "Returns live cash balance, portfolio value, and execution mode."],
        ["Service Beta (Broker)", "/api/v1/portfolio", "GET", "Analyzed holdings matrix, current market value, PnL, and recommendations."],
        ["Service Beta (Broker)", "/api/v1/place_order", "POST", "Pre-flight validation and order dispatch (Live Dhan API vs Paper log)."],
        ["Service Gamma (Intel)", "/api/v1/sentiment", "GET", "Overall market sentiment score, India VIX level, geopolitical risk."],
        ["Service Gamma (Intel)", "/api/v1/fii_dii", "GET", "FII and DII daily net equity buy/sell cash flow totals."],
        ["Service Gamma (Intel)", "/api/v1/fundamentals", "GET", "Stock fundamental financial metrics (P/E, P/B, Debt/Equity, ROE)."]
    ]

    for row_idx, row_data in enumerate(api_data, start=1):
        row_cells = t_api.rows[row_idx].cells
        bg_color = "F2F4F7" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = cell_value
            set_cell_background(row_cells[col_idx], bg_color)
            p = row_cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [1, 2] else WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.size = Pt(9)
                if col_idx == 0:
                    run.font.bold = True

    # Section 3: State-Graph Execution & Workflow Engine
    add_h1("3. LangGraph Workflow Execution Pipeline")
    doc.add_paragraph(
        "Trade execution is managed as a deterministic state graph powered by LangGraph. "
        "The workflow enforces strict sequential processing nodes to guarantee risk budgeting and regulatory compliance."
    )

    add_bullet("Node 1: Sync Portfolio Node -> ", "Connects to Google Sheets, syncs active positions, processes exits, and loads updated account funds.")
    add_bullet("Node 2: Screen Candidates Node -> ", "Executes strategy screener algorithms across target universes (Nifty 50, LargeMidcap 250, Champions Pool, ETFs).")
    add_bullet("Node 3: Risk Budgeting Node -> ", "Sizes positions using 1.5% portfolio risk and 2x ATR(14) stop loss. Enforces max 3 open positions per sector cap.")
    add_bullet("Node 4: Execute Orders Node -> ", "Dispatches validated orders to Service Beta Gateway for Dhan live API execution or Google Sheets paper logging.")
    add_bullet("Node 5: Format Scan Report Node -> ", "Formats trade execution report for Streamlit UI rendering and Telegram broadcast.")

    # Section 4: Data Storage & Google Sheets Schema
    add_h1("4. Persistence Layer & Google Sheets Architecture")
    doc.add_paragraph(
        "Persistent system state is saved in Google Sheets ('NSE_Dhan_Portfolio_Manager' / 'NSE_Nexus_Portfolio_Manager'). "
        "All gspread API operations utilize exponential backoff sleep retries to mitigate Google Sheets 429 rate limits."
    )

    add_bullet("1. Account Worksheet: ", "Stores Key-Value parameter pairs (Cash Balance, Total Portfolio Value, Risk Percent, Execution Mode).")
    add_bullet("2. OpenPositions Worksheet: ", "Tracks active holdings, entry date, quantity, entry price, initial SL, current SL, and target price.")
    add_bullet("3. TradeHistory Worksheet: ", "Permanent log of executed trades, exit prices, realized PnL, holding days, and exit reasons.")
    add_bullet("4. Paper_Trades Worksheet: ", "Simulated paper trade execution audit trail with timestamps and strategy rationales.")

    # Section 5: Security & Regulatory Governance
    add_h1("5. Security, Tax Sentinel & Regulatory Controls")
    doc.add_paragraph(
        "The system incorporates an autonomous Regulatory Sentinel module (tax_sentinel.py) and strict credential security controls:"
    )

    add_bullet("Individual Execution Mode Toggles: ", "Each strategy tab maintains an isolated Execution Mode radio control (PAPER_SIMULATED vs LIVE_DHAN) persisted in Google Sheets.")
    add_bullet("Pre-Flight Dhan OAuth Check: ", "Live order placement performs a mandatory pre-flight token check before dispatching orders.")
    add_bullet("Tax Sentinel & Transaction Fee Sizer: ", "Calculates Securities Transaction Tax (STT), Exchange Turnover Charges, GST, Stamp Duty, and Short-Term Capital Gains (STCG @ 20%).")

    output_path = r"C:\Users\ckbas\Documents\antigravity\AI-Swing-Trade\AI_Swing_Trade_System_Architecture.docx"
    artifact_path = r"C:\Users\ckbas\.gemini\antigravity\brain\393b1aff-76a5-403b-a924-df0af890188f\AI_Swing_Trade_System_Architecture.docx"

    doc.save(output_path)
    doc.save(artifact_path)
    print(f"Successfully generated System Architecture MS Word document at:\n{output_path}\n{artifact_path}")

if __name__ == "__main__":
    build_architecture_doc()
